#!/usr/bin/env python3
"""Build data/compilation_reconciliation.csv and data/year_audit_summary.csv

Two-way, row-level audit between FDA's official CDER Novel Drug Approvals
Compilation (data/raw/probe/fda_nme_compilation_1985_2025.xlsx, official Excel
published by FDA; SHA-256 manifest alongside) and this repository's master
list (data/fda_decisions_master.csv).

Forward direction: every one of the 1,387 Compilation rows (1985-2025) must be
represented in the master. Reverse direction: every master row dated <= 2025
must be explained relative to the Compilation (the expected master-only row is
D634 Contrave, kept for transparency and flagged NOT_ON_FDA_NME_TABLE).

Match cascade (deterministic, first hit wins, each hit is labelled):
  1. application number     - any of the row's 1-3 NDA/BLA numbers, extracted on
                              the master side from source_url_1 (ApplNo=) and
                              from the audited notes text ("application NNNNNN")
  2. brand + approval date  - exact normalised brand, same YYYY-MM-DD
  3. brand + year           - exact normalised brand in the same year
  4. brandnostop + date     - brand with '(' segments stripped both sides, same date
  5. generic + year         - normalised generic's first component matches the
                              master's generic's first component in the same year
                              (handles FDA's "[drug marketed without a proprietary
                              name]" convention vs the master's own naming)

Compilation XLSX schema (27 columns, verified 2026-09-17):
  0 Proprietary Name, 1 Active Ingredient/Moiety, 2 Applicant, 3 NDA/BLA,
  4-6 Application Number(1..3), 7-12 Dosage Form/Route(1..3),
  13 FDA Receipt Date, 14 FDA Approval Date, 15 Approval Year,
  16 Abbreviated Indication(s), 17 Approved Use(s), 18 Review Designation,
  19 Orphan, 20 Accelerated Approval, 21 Breakthrough, 22 Fast Track,
  23 QIDP, 24 PRV issued, 25 PRV redeemed, 26 Notes.

Known by-design exclusions (documented, surfaced in the outputs):
  * Emend IV (fosaprepitant, NDA 022023, 2008) is on the Compilation but FDA's
    contemporaneous 2008 NME table excluded it (aprepitant moiety approved as
    Emend in 2003). The master follows FDA's year table — the row is reported
    as NOT_IN_MASTER with that documented reason, never silently dropped.

Run: .venv/bin/python scripts/build_compilation_reconciliation.py
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "raw" / "probe" / "fda_nme_compilation_1985_2025.xlsx"
MASTER = ROOT / "data" / "fda_decisions_master.csv"
CROSSCHECK = ROOT / "data" / "verification_crosscheck.csv"
OUT_RECON = ROOT / "data" / "compilation_reconciliation.csv"
OUT_YEAR = ROOT / "data" / "year_audit_summary.csv"

COMPILATION_SOURCE = "https://www.fda.gov/media/177921/download?attachment"

# The one documented Compilation row the master deliberately excludes.
BY_DESIGN = {"NDA22023": "Compilation row for Emend IV (fosaprepitant, NDA 022023, 2008) "
             "excluded by design: FDA's own 2008 NME table excludes fosaprepitant "
             "(aprepitant moiety already approved 2003 as Emend, D985 master row). "
             "Surfaced here, never silently patched."}


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def strip_parens(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\([^)]*\)", " ", str(s or ""))).strip().lower()


# ---------------------------------------------------------------- compilation
def load_compilation() -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    out = []
    for r in rows[1:]:
        if r[0] is None and r[1] is None:
            continue
        if str(r[0] or "").startswith("*"):
            continue
        if not isinstance(r[15], int):
            continue
        def ds(x):
            return x.strftime("%Y-%m-%d") if hasattr(x, "strftime") else ("" if x is None else str(x)[:10])
        appl_nums = [str(int(a)) for a in (r[4], r[5], r[6])
                     if a is not None and str(a).strip() not in {"", "None"}]
        out.append({
            "year": r[15],
            "brand": (r[0] or "").strip(),
            "generic": (r[1] or "").strip(),
            "applicant": (r[2] or "").strip(),
            "kind": (r[3] or "").strip(),
            "appl_numbers": appl_nums,
            "applications": [f"{(r[3] or '').strip()}{n}" for n in appl_nums],
            "approval_date": ds(r[14]),
            "abbrev_ind": (r[16] or "").strip() if r[16] else "",
            "review": (r[18] or "").strip() if r[18] else "",
            "accelerated": (r[20] or "").strip() if r[20] else "",
        })
    return out


# -------------------------------------------------------------------- master
APPL_PATTERNS = [
    re.compile(r"(?:varApplNo|ApplNo)=\s*0*(\d{5,7})", re.I),
    re.compile(r"application\s+(?:number\s+)?0*(\d{5,7})", re.I),
    re.compile(r"\b(?:NDA|BLA)\s*0?#?\s*0*(\d{5,7})\b", re.I),
    re.compile(r"\bN0*(\d{5,7})\b"),  # 'appl N020831' texts
]


def master_rows() -> list[dict]:
    rows = list(csv.DictReader(open(MASTER, newline="", encoding="utf-8-sig")))
    for r in rows:
        nums = set()
        for src in (r.get("source_url_1", ""), r.get("source_url_2", ""), r.get("notes", "")):
            for pat in APPL_PATTERNS:
                for mo in pat.finditer(src or ""):
                    nums.add(mo.group(1).lstrip("0") or "0")
        r["_appls"] = nums
        r["_year"] = (r.get("decision_date") or "")[:4]
        b = r.get("drug_brand", "")
        r["_bn"] = norm(b)
        r["_bns"] = norm(strip_parens(b))
        g = re.split(r"[;,]", r.get("drug_generic", ""))[0]
        r["_g0"] = norm(re.sub(r"\s+", " ", g).strip())[:30]
    return rows


def load_crosscheck() -> dict:
    out = {}
    try:
        for r in csv.DictReader(open(CROSSCHECK, newline="", encoding="utf-8-sig")):
            out[r["decision_id"]] = r.get("crosscheck_status", "")
    except FileNotFoundError:
        pass
    return out


def build_indexes(rows: list[dict]):
    by_appl = defaultdict(list)
    by_bd = defaultdict(list)
    by_by = defaultdict(list)
    by_bsd = defaultdict(list)
    by_gy = defaultdict(list)
    for r in rows:
        for a in r["_appls"]:
            by_appl[a].append(r)
        by_bd[(r["_bn"], r["decision_date"])].append(r)
        by_by[(r["_bn"], r["_year"])].append(r)
        by_bsd[(r["_bns"], r["decision_date"])].append(r)
        by_gy[(r["_g0"], r["_year"])].append(r)
    return by_appl, by_bd, by_by, by_bsd, by_gy


def flag_of(r: dict) -> str:
    """Row-level flag taxonomy. NOT_ON_FDA_NME_TABLE detection checks
    decision_type + verification_status + notes, but strips attribution text
    ('D634 re-labelled NOT_ON_FDA_NME_TABLE' inside Ofev's note must NOT flag
    Ofev itself)."""
    n = r.get("decision_type", "") + " " + r.get("verification_status", "") + " " + r.get("notes", "")
    n = re.sub(r"D\d+\s+re-?labelled\s+NOT_ON_FDA_NME_TABLE", "", n, flags=re.I)
    if "NOT_ON_FDA_NME_TABLE" in n:
        return "NOT_ON_FDA_NME_TABLE"
    n = r.get("notes", "")
    if "COMPILATION_ONLY_CBER" in n:
        return "COMPILATION_ONLY_CBER"
    if "COMPILATION_ONLY" in n:
        return "COMPILATION_ONLY"
    return "YEAR_TABLE_ROW"


def main() -> int:
    comp = load_compilation()
    master = master_rows()
    cross = load_crosscheck()
    by_appl, by_bd, by_by, by_bsd, by_gy = build_indexes(master)

    recon = []
    matched_ids = set()
    yeargrid = defaultdict(lambda: Counter())
    for c in comp:
        y = str(c["year"])
        yeargrid[y]["compilation_rows"] += 1
        found, via = None, ""
        for a in c["appl_numbers"]:
            hits = by_appl.get((a.lstrip("0") or "0"), [])
            if hits:
                found, via = hits[0], f"appl {c['kind']}{a}"
                break
        if not found:
            hits = by_bd.get((norm(c["brand"]), c["approval_date"]) or "")
            if hits:
                found, via = hits[0], "brand+date"
        if not found:
            hits = by_by.get((norm(c["brand"]), y))
            if hits:
                found, via = hits[0], "brand+year"
        if not found:
            hits = by_bsd.get((norm(strip_parens(c["brand"])), c["approval_date"]))
            if hits:
                found, via = hits[0], "brand-no-qualifier+date"
        if not found:
            g0 = norm(c["generic"].split(";")[0].split(",")[0])[:30]
            if g0 and len(g0) >= 6:
                hits = [r for (g, yy), rs in by_gy.items() if yy == y and g == g0 for r in rs] or \
                       [r for (g, yy), rs in by_gy.items() if yy == y and (g.startswith(g0[:20]) or g0.startswith(g[:20])) for r in rs]
                if hits:
                    found, via = hits[0], "generic+year"
        notes = ""
        status = "MATCHED"
        if found:
            matched_ids.add(found["decision_id"])
            yeargrid[y]["matched"] += 1
            did, flag = found["decision_id"], flag_of(found)
        else:
            did, flag = "", ""
            key = f"{c['kind']}{c['appl_numbers'][0]}" if c["appl_numbers"] else c["brand"]
            if key in BY_DESIGN:
                status = "NOT_IN_MASTER"
                notes = BY_DESIGN[key]
                yeargrid[y]["missing_by_design"] += 1
            else:
                status = "NOT_IN_MASTER"
                yeargrid[y]["missing"] += 1
        recon.append({
            "direction": "compilation→master",
            "approval_year": y,
            "brand": c["brand"],
            "generic": c["generic"],
            "applicant": c["applicant"],
            "application": ";".join(c["applications"]),
            "approval_date": c["approval_date"],
            "compilation_review": c["review"],
            "compilation_accelerated": c["accelerated"],
            "reconciliation_status": status,
            "matched_via": via,
            "master_decision_id": did,
            "master_flag": flag,
            "openfda_crosscheck": cross.get(did, "") if did else "",
            "compilation_source": COMPILATION_SOURCE + ("" if not notes else "  |  NOTE: " + notes),
        })

    # reverse direction: master rows that the Compilation does not carry
    master_only = []
    for r in master:
        if r["decision_id"] in matched_ids:
            continue
        y = r["_year"]
        yeargrid[y]["master_rows"] += 1
        yeargrid[y]["missing"] += 0
        status = "NOT_ON_COMPILATION"
        if int(y or 0) >= 2026:
            status = "2026_DECISION_BEYOND_COMPILATION_EDITION"
        elif flag_of(r) == "NOT_ON_FDA_NME_TABLE":
            status = "NOT_ON_COMPILATION_DOCUMENTED"
        master_only.append({
            "direction": "master→compilation",
            "approval_year": y,
            "brand": r["drug_brand"],
            "generic": r["drug_generic"],
            "applicant": r["company_name"],
            "application": ";".join(sorted(r["_appls"])),
            "approval_date": r["decision_date"],
            "compilation_review": "",
            "compilation_accelerated": "",
            "reconciliation_status": status,
            "matched_via": "",
            "master_decision_id": r["decision_id"],
            "master_flag": flag_of(r),
            "openfda_crosscheck": cross.get(r["decision_id"], ""),
            "compilation_source": "" if not status.startswith("NOT_ON_COMPILATION_DOC") else
            "Currently on FDA year-table enumeration only (see row notes; NOT_ON_FDA_NME_TABLE flag).",
        })

    # fill forward master_rows counts for all years
    for r in master:
        yeargrid[r["_year"]]["master_rows_total"] += 1

    allrows = recon + master_only
    allrows.sort(key=lambda x: (str(x["approval_year"]), x["direction"], x["brand"]))
    fields = ["direction", "approval_year", "brand", "generic", "applicant", "application",
              "approval_date", "compilation_review", "compilation_accelerated",
              "reconciliation_status", "matched_via", "master_decision_id", "master_flag",
              "openfda_crosscheck", "compilation_source"]
    with open(OUT_RECON, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader(); w.writerows(allrows)

    with open(OUT_YEAR, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["year", "compilation_rows", "matched", "missing", "missing_by_design",
                    "master_rows_total", "master_only_rows", "compilation_2000_2003_cber",
                    "master_flag_COMPILATION_ONLY", "master_flag_COMPILATION_ONLY_CBER",
                    "master_flag_NOT_ON_FDA_NME_TABLE"])
        for y in sorted(yeargrid):
            c = yeargrid[y]
            yrs = [r for r in master if r["_year"] == y]
            w.writerow([y, c["compilation_rows"], c["matched"], c["missing"], c["missing_by_design"],
                        c["master_rows_total"],
                        sum(1 for mo in master_only if mo["approval_year"] == y),
                        sum(1 for r in yrs if flag_of(r) == "COMPILATION_ONLY_CBER" and y in {"2000", "2001", "2002", "2003"}),
                        sum(1 for r in yrs if flag_of(r) == "COMPILATION_ONLY"),
                        sum(1 for r in yrs if flag_of(r) == "COMPILATION_ONLY_CBER"),
                        sum(1 for r in yrs if flag_of(r) == "NOT_ON_FDA_NME_TABLE"),
                        ])

    n_matched = sum(1 for r in recon if r["reconciliation_status"] == "MATCHED")
    n_missing = [r for r in recon if r["reconciliation_status"] != "MATCHED"]
    print(f"compilation rows: {len(recon)} | matched in master: {n_matched} | not matched: {len(n_missing)}")
    for r in n_missing:
        print("  NOT_MATCHED:", r["approval_year"], r["brand"], r["application"],
              "(by design)" if "NOTE:" in r["compilation_source"] else "(UNEXPLAINED)")
    print(f"master-only rows: {len(master_only)}")
    tall = Counter((r["approval_year"], r["reconciliation_status"]) for r in master_only)
    for k in sorted(tall):
        print("  ", k, tall[k])
    print(f"wrote {OUT_RECON.relative_to(ROOT)} ({len(allrows)} rows incl. header-less)")
    print(f"wrote {OUT_YEAR.relative_to(ROOT)}")
    bad = [r for r in n_missing if "NOTE:" not in r["compilation_source"]]
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
